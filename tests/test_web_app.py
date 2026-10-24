"""User-visible CV review, real-catalog integration, and session boundaries."""

import unittest
import re
import os
from unittest.mock import patch
from io import BytesIO

from fastapi.testclient import TestClient

from src.main import app
from src.profiles.cv_parser import MAX_BYTES, CVParseError, parse_cv
from src.profiles.samples import SAMPLE_PATH, sample_cv_pdf

ROLE = "role:onet_15_2051_01"


class TestWebApp(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        environment = patch.dict(os.environ, {"TREND_AUTO_REFRESH": "0", "GROQ_API": "", "GROQ_API_KEY": "", "GEMINI_API_KEY": "", "RUNPOD_API_KEY": "", "RUNPOD_ENDPOINT_ID": ""})
        environment.start()
        cls.addClassCleanup(environment.stop)
        cls.client = TestClient(app)
        cls.client.__enter__()

    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None, None, None)

    def upload(self, text=b"Skills: Python, SQL"):
        response = self.client.post("/api/v2/profiles/upload", content=text,
                                    headers={"X-Filename": "fictional.txt"})
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        return data, {"X-Session-ID": data["session_id"]}

    def review(self, data, headers, **extra):
        payload = {"revision": data["revision"],
                   "skills": [{**skill, "confirmation": "confirmed"}
                              for skill in data["profile"]["skills"]], **extra}
        response = self.client.patch("/api/v2/profiles/current", json=payload, headers=headers)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_dependency_map_requires_review_and_respects_session_revision(self):
        data, headers = self.upload()
        payload = {"revision": data["revision"], "role_id": ROLE, "horizon": "short", "weekly_hours": 5}
        self.assertEqual(self.client.post("/api/v2/pathways/dependencies", json=payload, headers=headers).status_code, 409)
        reviewed = self.review(data, headers)
        payload["revision"] = reviewed["revision"]
        response = self.client.post("/api/v2/pathways/dependencies", json=payload, headers=headers)
        self.assertEqual(response.status_code, 200, response.text)
        graph = response.json()["roadmap"]
        self.assertEqual(graph["origin"], "application_outline")
        self.assertEqual(graph["status"], "not_configured")
        self.assertLessEqual(graph["planned_hours"], 60)
        self.assertLessEqual(len(graph["nodes"]), 5)
        self.assertFalse(graph["changes_personal_fit"])
        for invalid in ({"horizon": "forever"}, {"weekly_hours": 0}, {"weekly_hours": 169},
                        {"selected_software_skill_ids": ["skill:invented"]}):
            self.assertEqual(self.client.post("/api/v2/pathways/dependencies", json={**payload, **invalid}, headers=headers).status_code, 422)
        payload["revision"] -= 1
        self.assertEqual(self.client.post("/api/v2/pathways/dependencies", json=payload, headers=headers).status_code, 409)

    def test_dependency_ai_context_omits_cv_evidence_and_identifiers(self):
        from src.reskilling.dependencies import application_outline, validate_proposal
        data, headers = self.upload()
        reviewed = self.review(data, headers)
        captured = []
        def inspect(context):
            captured.append(context)
            return validate_proposal(application_outline(context), context, origin="application_outline")
        with patch.object(app.state.dependency_maps, "curate", side_effect=inspect):
            response = self.client.post("/api/v2/pathways/dependencies", headers=headers, json={
                "revision": reviewed["revision"], "role_id": ROLE, "horizon": "standard", "weekly_hours": 5})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertNotIn("profile_id", captured[0])
        self.assertNotIn("session_id", captured[0])
        for skill in captured[0]["known_skills"]:
            self.assertEqual(set(skill), {"skill_id", "label", "level"})

    def test_language_options_validate_and_fallback_without_changing_learning_facts(self):
        data, headers = self.upload()
        reviewed = self.review(data, headers)
        payload = {"revision": reviewed["revision"], "role_id": ROLE, "weekly_hours": 5}
        english = self.client.post('/api/v2/pathways/dependencies', headers=headers, json=payload).json()['roadmap']
        for language in ('ms', 'zh'):
            response = self.client.post('/api/v2/pathways/dependencies', headers=headers, json={**payload, 'language': language})
            self.assertEqual(response.status_code, 200, response.text)
            result = response.json()['roadmap']
            self.assertEqual(result['nodes'], english['nodes'])
            self.assertEqual(result['edges'], english['edges'])
            self.assertEqual(result['content_language'], 'en')
            self.assertEqual(result['translation_status'], 'unavailable')
        self.assertEqual(self.client.post('/api/v2/pathways/dependencies', headers=headers, json={**payload, 'language': 'fr'}).status_code, 422)
        self.assertEqual(self.client.get(f'/api/v2/roles/{ROLE}/trends?language=fr', headers=headers).status_code, 422)
        self.assertEqual(self.client.get(f'/api/v2/roles/{ROLE}/trends?language=ms', headers=headers).status_code, 200)

    def test_python_level_edits_refresh_matching_score(self):
        data, headers = self.upload(b"Skills: Python")
        scores = []
        for level in ("beginner", "intermediate", "advanced"):
            skills = [{**skill, "level": level, "confirmation": "confirmed"} for skill in data["profile"]["skills"]]
            response = self.client.patch("/api/v2/profiles/current", headers=headers,
                json={"revision": data["revision"], "skills": skills, "desired_role_id": ROLE})
            self.assertEqual(response.status_code, 200, response.text)
            data = response.json()
            response = self.client.get("/api/v2/matches", headers=headers, params={"revision": data["revision"]})
            self.assertEqual(response.status_code, 200, response.text)
            row = next(row for row in response.json()["matches"] if row["role_id"] == ROLE)
            self.assertEqual(row["score_breakdown"]["alignment_contributions"][0]["profile_level"], level)
            scores.append(row["fit_score"])
        # This occupation has an estimated intermediate Python requirement.
        self.assertEqual(scores, [6.5, 16.25, 12.88])

    def test_rina_upload_exceeds_ninety_with_ai_levels_and_confirmation_only(self):
        from scripts.create_high_match_cv import OUTPUT, ROLE as MARKETING
        from src.profiles.assessment import AssessmentBatch
        from types import SimpleNamespace
        import json
        from unittest.mock import Mock
        levels = {s["skill_id"]: s["level"] for s in json.loads((OUTPUT / "reviewed_profile.json").read_text())["skills"]}
        def generate(instruction, data, response_model, **kwargs):
            self.assertEqual(len(data["skills"]), 39)
            rows = [dict(skill_id=s["skill_id"], level=levels[s["skill_id"]], confidence="medium",
                rationale="Supported by independent task evidence.", evidence_indices=[s["evidence"][0]["index"]]) for s in data["skills"]]
            return SimpleNamespace(value=AssessmentBatch(skills=rows), returned_model="test-complete", provider="groq")
        with patch("src.profiles.assessment.assessment_configuration", return_value=("groq", [])), \
             patch("src.profiles.assessment.assessment_client", return_value=Mock(generate=generate)):
            response = self.client.post("/api/v2/profiles/upload", content=(OUTPUT / "rina_search_marketing.pdf").read_bytes(), headers={"X-Filename": "rina.pdf"})
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()
        self.assertEqual(data["profile"]["assessment_status"]["assessed_skill_count"], 39)
        reviewed = self.review(data, {"X-Session-ID": data["session_id"]}, desired_role_id=MARKETING)
        rows = self.client.get("/api/v2/matches", headers={"X-Session-ID": data["session_id"]}, params={"revision": reviewed["revision"]}).json()["matches"]
        match = next(row for row in rows if row["role_id"] == MARKETING)
        self.assertGreater(match["fit_score"], 90)
        self.assertEqual(match["rank"], 1)

    def test_live_trend_api_serves_fetched_sources_and_protects_refresh(self):
        import json
        from tempfile import TemporaryDirectory
        from pathlib import Path
        from src.intelligence.live_trends import LiveTrendService
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            config = root / "feeds.json"
            config.write_text(json.dumps({"schema_version": "0.1.0", "lookback_days": 90,
                "max_items_per_feed": 10, "keywords": ["AI"], "feeds": [{"feed_id": "test",
                "publisher": "TEST source", "url": "https://research.example.org/feed",
                "article_prefix": "https://research.example.org/articles/"}]}))
            from datetime import datetime, UTC
            body = b"<rss><channel><item><title>AI research fixture</title><link>https://research.example.org/articles/one</link><pubDate>Thu, 22 Oct 2026 16:00:00 +0000</pubDate><description>AI agents can assist report preparation while people check the outputs and original data sources.</description></item></channel></rss>"
            service = LiveTrendService(app.state.backend.catalog, path=root / "cache.json", seed=root / "none.json",
                config_path=config, fetcher=lambda url: body, article_fetcher=lambda url, prefix: "AI research fixture supports report preparation and checking original source records. " * 2, now=lambda: datetime(2026, 10, 23, 12, tzinfo=UTC))
            original = app.state.trends
            app.state.trends = service
            try:
                service.refresh()
                response = self.client.get(f"/api/v2/roles/{ROLE}/trends")
                self.assertEqual(response.status_code, 200)
                data = response.json()
                self.assertEqual(data["articles"][0]["published_at"], "2026-10-22")
                self.assertEqual(data["refresh"]["status"], "fresh")
                self.assertEqual(data["source_review"]["stored_count"], 1)
                self.assertEqual(data["source_review"]["assessed_count"], 0)
                self.assertFalse(data["source_review"]["articles"][0]["assessed_by_ai"])
                self.assertFalse(data["changes_personal_fit"])
                self.assertNotIn("source_text", data["articles"][0])
                self.assertEqual(self.client.post(f"/api/v2/roles/{ROLE}/trends/refresh").status_code, 422)
                session = self.client.post("/api/v2/profiles/manual").json()["session_id"]
                with patch.object(service, "start_refresh", return_value=False):
                    response = self.client.post(f"/api/v2/roles/{ROLE}/trends/refresh", headers={"X-Session-ID": session})
                self.assertEqual(response.status_code, 202)
                self.assertEqual(response.json()["articles"][0]["title"], "AI research fixture")
                self.assertEqual(self.client.post(f"/api/v2/roles/{ROLE}/trends/retry").status_code, 422)
                with patch.object(service, "ensure_analysis") as analyze, patch.object(service, "start_refresh") as fetch:
                    response = self.client.post(f"/api/v2/roles/{ROLE}/trends/retry", headers={"X-Session-ID": session})
                self.assertEqual(response.status_code, 202)
                analyze.assert_called_once_with(ROLE)
                fetch.assert_not_called()
                self.assertEqual(self.client.get("/api/v2/roles/role:unknown/trends").status_code, 404)
            finally:
                app.state.trends = original
                service.close()

    def test_future_scenarios_are_complete_offline_and_updates_require_a_session(self):
        response = self.client.get(f"/api/v2/roles/{ROLE}/futures")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn(data["origin"], {"application_authored", "ai_generated"})
        self.assertEqual({r["scenario_id"] for r in data["stories"]}, {"moderate_ai_change", "rapid_ai_change"})
        self.assertTrue(data["skill_comparison"])
        self.assertFalse(data["changes_personal_fit"])
        self.assertTrue(data["stories"][0]["task_changes"][0]["human_responsibility"])
        self.assertEqual(self.client.get("/api/v2/roles/unknown/futures").status_code, 404)
        self.assertEqual(self.client.post(f"/api/v2/roles/{ROLE}/futures/refresh").status_code, 422)
        session = self.client.post("/api/v2/profiles/manual").json()["session_id"]
        with patch.object(app.state.futures, "ensure") as ensure:
            response = self.client.post(f"/api/v2/roles/{ROLE}/futures/refresh", headers={"X-Session-ID": session})
        self.assertEqual(response.status_code, 202)
        ensure.assert_called_once_with(ROLE)
        self.assertEqual(len(response.json()["stories"]), 2)

    def test_shell_assets_and_downloadable_cv(self):
        page = self.client.get("/")
        self.assertEqual(page.status_code, 200)
        self.assertIn('id="root"', page.text)
        self.assertIn("script-src 'self'", page.headers["Content-Security-Policy"])
        assets = re.findall(r'(?:src|href)="(/assets/[^"]+)"', page.text)
        self.assertTrue(any(path.endswith(".js") for path in assets))
        self.assertTrue(any(path.endswith(".css") for path in assets))
        for path in [*assets, "/fonts/InterVariable.woff2", "/health"]:
            self.assertEqual(self.client.get(path).status_code, 200)
        pdf = self.client.get("/api/v2/samples/graduate.pdf")
        self.assertEqual(pdf.headers["content-type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF-"))
        self.assertEqual(SAMPLE_PATH.with_suffix(".pdf").read_bytes(), sample_cv_pdf())
        # A text extractor can tolerate malformed page streams that viewers reject.
        # Check the actual downloaded bytes with the independent PDFium renderer.
        import pypdfium2 as pdfium

        with pdfium.PdfDocument(pdf.content) as document:
            page = document[0]
            text = page.get_textpage()
            self.assertIn("AISHA TAN", text.get_text_range())
            text.close()
            bitmap = page.render(scale=1)
            image = bitmap.to_pil().convert("L")
            self.assertLess(image.getextrema()[0], 100, "Downloaded CV rendered as a blank page")
            self.assertGreater(sum(image.histogram()[:200]), 1000)
            bitmap.close()
            page.close()
        parsed = parse_cv(pdf.content, "sample.pdf")
        self.assertIn("AISHA TAN", parsed.pages[0])
        sample = self.client.post("/api/v2/profiles/sample").json()
        uploaded = self.client.post("/api/v2/profiles/upload", content=pdf.content,
                                   headers={"X-Filename": "downloaded.pdf"}).json()
        self.assertEqual(sample["profile"]["skills"], uploaded["profile"]["skills"])
        self.assertTrue(all(row["confirmation"] == "unverified"
                            for row in uploaded["profile"]["skills"]))

    def test_review_required_and_input_changes_computed_matches(self):
        data, headers = self.upload()
        self.assertEqual(self.client.get("/api/v2/matches?revision=0", headers=headers).status_code, 409)
        reviewed = self.review(data, headers)
        first = self.client.get("/api/v2/matches?revision=1", headers=headers).json()["matches"]
        self.assertTrue(first)
        self.assertEqual(first[0]["schema_version"], "0.2.0")
        reviewed = self.review(reviewed, headers, skills=[])
        empty = self.client.get("/api/v2/matches?revision=2", headers=headers).json()["matches"]
        self.assertTrue(all(row["fit_score"] == 0 for row in empty))
        self.assertNotEqual(first, empty)

    def test_ai_retry_applies_levels_and_rejects_stale_requests(self):
        from types import SimpleNamespace
        from src.profiles.assessment import AssessmentBatch

        data, headers = self.upload(b"Used Python independently to build a fictional dashboard.")
        skill_id = data["profile"]["skills"][0]["skill_id"]
        result = SimpleNamespace(returned_model="test-model", value=AssessmentBatch(skills=[{
            "skill_id": skill_id, "level": "intermediate", "confidence": "medium",
            "rationale": "Independent project application.", "evidence_indices": [0]}]))
        with patch("src.profiles.assessment.assessment_configuration", return_value=("runpod", [])), \
             patch("src.profiles.assessment.assessment_client") as factory:
            factory.return_value.generate.return_value = result
            response = self.client.post("/api/v2/profiles/assess", headers=headers,
                json={"revision": 0, "skills": data["profile"]["skills"]})
        self.assertEqual(response.status_code, 200, response.text)
        assessed = response.json()
        self.assertEqual(assessed["profile"]["skills"][0]["level"], "intermediate")
        self.assertEqual(assessed["profile"]["assessment_status"]["provider"], "runpod")
        self.assertEqual(assessed["profile"]["skills"][0]["confirmation"], "unverified")
        self.assertEqual(self.client.post("/api/v2/profiles/assess", headers=headers,
            json={"revision": 0, "skills": data["profile"]["skills"]}).status_code, 409)

    def test_actual_gaps_scenarios_and_fit_preservation(self):
        data, headers = self.upload()
        reviewed = self.review(data, headers, desired_role_id=ROLE,
                               constraints={"budget_myr": 0, "hours_per_week": 5})
        matches = self.client.get("/api/v2/matches?revision=1", headers=headers).json()["matches"]
        selected = next(row for row in matches if row["role_id"] == ROLE)
        request = {"revision": reviewed["revision"], "role_id": ROLE}
        baseline = self.client.post("/api/v2/pathways", headers=headers, json=request)
        self.assertEqual(baseline.status_code, 200, baseline.text)
        plan = baseline.json()["plan"]
        sdg = plan["sdg_summary"]
        self.assertTrue(sdg["role_connections"])
        self.assertEqual(sdg["mapped_skill_count"], sdg["skill_count"])
        self.assertTrue(all(s["connections"] for s in sdg["skills"]))
        self.assertTrue(any(s["status"] == "confirmed" for s in sdg["skills"]))
        self.assertTrue(all("skill:" not in step["learning_goal"] for m in plan["milestones"] for step in m["steps"]))
        rapid = self.client.post("/api/v2/pathways", headers=headers,
                                 json={**request, "scenario_id": "rapid_ai_change"}).json()["plan"]
        self.assertEqual(plan["fit_score"], selected["fit_score"])
        self.assertEqual(plan["fit_score"], rapid["fit_score"])
        self.assertNotEqual(plan["prioritized_gaps"], rapid["prioritized_gaps"])
        restored = self.client.post("/api/v2/pathways", headers=headers, json=request).json()["plan"]
        self.assertEqual(plan["prioritized_gaps"], restored["prioritized_gaps"])
        self.assertTrue(plan["selected_courses"])
        self.assertTrue(all(course["cost"]["amount"] == 0 for course in plan["selected_courses"]))
        self.assertTrue(plan["sdg_connections"])
        self.assertTrue(any(step["resource_status"] == "verified"
                            for milestone in plan["milestones"] for step in milestone["steps"]))

    def test_unsupported_scenarios_and_unknown_roles(self):
        roles = self.client.get("/api/v2/roles?limit=2").json()
        self.assertEqual(roles["total"], 910)
        role_id = roles["roles"][0]["role_id"]
        self.assertNotEqual(role_id, ROLE)
        baseline = self.client.get(f"/api/v2/roles/{role_id}/outlook").json()
        self.assertEqual(baseline["available_scenarios"], ["baseline", "moderate_ai_change", "rapid_ai_change"])
        self.assertEqual(baseline["baseline_label"], "Insufficient evidence")
        response = self.client.get(f"/api/v2/roles/{role_id}/outlook?scenario_id=rapid_ai_change")
        self.assertEqual(response.status_code, 200)
        rapid = response.json()
        self.assertNotEqual(baseline["scenario"]["skill_priorities"], rapid["scenario"]["skill_priorities"])
        self.assertEqual(baseline["sources"], rapid["sources"])
        # Current research comes from /trends, rather than a static fallback
        # attached to every occupation's reviewed outlook.
        self.assertEqual(baseline["general_research"], [])
        self.assertEqual(self.client.get("/api/v2/roles/unknown").status_code, 404)
        self.assertEqual(self.client.get(f"/api/v2/roles/{role_id}/evidence").json()["chains"], [])

    def test_new_evidence_routes_are_offline_and_expose_only_approved_results(self):
        with patch("src.intelligence.gemini.GeminiClient.generate", side_effect=AssertionError("Public requests must never call Gemini")):
            skills = self.client.get(f"/api/v2/roles/{ROLE}/skill-evidence")
            impacts = self.client.get(f"/api/v2/roles/{ROLE}/impacts")
        self.assertEqual(skills.status_code, 200)
        self.assertEqual(impacts.status_code, 200)
        self.assertEqual(skills.json()["approved_conditional_skill_evidence"], [])
        self.assertEqual(impacts.json()["approved_conditional_role_impacts"], [])
        self.assertFalse(skills.json()["changes_weights"])
        self.assertFalse(impacts.json()["changes_weights"])
        with patch.object(self.client.app.state.backend, "skill_evidence", return_value={
            **skills.json(), "approved_conditional_skill_evidence": [{"skill_id": "skill:sql", "rationale": "Fictional test interpretation"}],
        }):
            result = self.client.get(f"/api/v2/roles/{ROLE}/skill-evidence").json()
        self.assertEqual(result["approved_conditional_skill_evidence"][0]["skill_label"], "SQL")
        for suffix in ("skill-evidence", "impacts"):
            self.assertEqual(self.client.get(f"/api/v2/roles/unknown/{suffix}").status_code, 404)
        with patch.object(self.client.app.state.backend, "skill_evidence", side_effect=ValueError("private backend detail")):
            result = self.client.get(f"/api/v2/roles/{ROLE}/skill-evidence")
        self.assertEqual(result.status_code, 503)
        self.assertNotIn("private backend detail", result.text)

    def test_shared_course_store_is_read_by_status_and_planner(self):
        from tempfile import TemporaryDirectory
        from pathlib import Path
        from src.reskilling.course_catalog import CourseCatalogStore

        backend = self.client.app.state.backend
        original = backend.course_store
        with TemporaryDirectory() as directory:
            store = CourseCatalogStore(Path(directory) / "courses.json")
            backend.course_store = store
            try:
                response = self.client.get("/api/v2/courses/status")
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()["status"], "empty")
                self.assertIsNone(response.json()["updated_at"])
                data, headers = self.upload()
                reviewed = self.review(data, headers, desired_role_id=ROLE)
                matches = self.client.get(f"/api/v2/matches?revision={reviewed['revision']}", headers=headers).json()["matches"]
                gap = next(row for row in matches if row["role_id"] == ROLE)["missing_skills"][0]
                from datetime import UTC, datetime
                from src.schemas.roadmap import VerifiedCourse
                store.upsert(VerifiedCourse(
                    course_id="course:fictional_integration", title="Fictional integration course",
                    provider="Fictional test provider", course_url="https://example.org/test-course",
                    evidence_url="https://example.org/test-course", source_tier="trusted",
                    skill_ids=[gap["skill_id"]], coverage_strength={gap["skill_id"]: 1},
                    cost={"access": "free", "amount": 0, "includes_certificate": False},
                    duration={"hours": 1, "display_text": "1 hour", "basis": "provider_stated"},
                    delivery_mode="online", language="English", last_verified_at=datetime.now(UTC)))
                self.assertEqual(self.client.get("/api/v2/courses/status").json()["verified_course_count"], 1)
                result = self.client.post("/api/v2/pathways", headers=headers, json={"revision": reviewed["revision"], "role_id": ROLE})
                self.assertEqual(result.status_code, 200, result.text)
                self.assertIn("course:fictional_integration", [course["course_id"] for course in result.json()["plan"]["selected_courses"]])
                store.path.write_text("invalid JSON", encoding="utf-8")
                self.assertEqual(self.client.get("/api/v2/courses/status").status_code, 503)
            finally:
                backend.course_store = original

    def test_sessions_stale_revisions_and_forgetting(self):
        data, headers = self.upload()
        self.review(data, headers)
        self.assertEqual(self.client.get("/api/v2/matches?revision=0", headers=headers).status_code, 409)
        self.assertEqual(self.client.get("/api/v2/matches?revision=1", headers={"X-Session-ID": "wrong"}).status_code, 401)
        self.assertEqual(self.client.patch("/api/v2/profiles/current", headers=headers,
                         json={"revision": 0, "skills": []}).status_code, 409)
        self.assertEqual(self.client.delete("/api/v2/profiles/current", headers=headers).status_code, 204)
        self.assertEqual(self.client.get("/api/v2/matches?revision=1", headers=headers).status_code, 401)

    def test_upload_bounds_and_private_validation_errors(self):
        oversized = self.client.post("/api/v2/profiles/upload", content=b"a" * (MAX_BYTES + 1))
        self.assertEqual(oversized.status_code, 413)
        malformed = self.client.post("/api/v2/profiles/upload", content=b"private cv text",
                                     headers={"X-Filename": "wrong.pdf"})
        self.assertEqual(malformed.status_code, 422)
        self.assertNotIn("private cv text", malformed.text)
        data, headers = self.upload()
        invalid = self.client.patch("/api/v2/profiles/current", headers=headers, json={
            "revision": 0, "skills": [{"evidence": ["DO_NOT_ECHO_THIS"]}]})
        self.assertEqual(invalid.status_code, 422)
        self.assertNotIn("DO_NOT_ECHO_THIS", invalid.text)
        self.assertEqual(invalid.headers["Cache-Control"], "no-store")
        expired = app.state.sessions.get(data["session_id"])
        expired.expires_at = 0
        self.assertEqual(self.client.get("/api/v2/matches?revision=0", headers=headers).status_code, 401)

    def test_review_rejects_invented_skills_and_unconfirmed_claims(self):
        data, headers = self.upload()
        unconfirmed = self.client.patch("/api/v2/profiles/current", headers=headers,
            json={"revision": 0, "skills": data["profile"]["skills"]})
        self.assertEqual(unconfirmed.status_code, 422)
        skill = {**data["profile"]["skills"][0], "label": "Invented", "confirmation": "confirmed"}
        self.assertEqual(self.client.patch("/api/v2/profiles/current", headers=headers,
            json={"revision": 0, "skills": [skill]}).status_code, 422)

    def test_optional_software_is_explicit_and_validated(self):
        data, headers = self.upload()
        self.review(data, headers, desired_role_id=ROLE)
        matches = self.client.get("/api/v2/matches?revision=1", headers=headers).json()["matches"]
        role = next(row for row in matches if row["role_id"] == ROLE)
        selected = role["software_learning_options"][0]["skill_id"]
        request = {"revision": 1, "role_id": ROLE, "selected_software_skill_ids": [selected]}
        response = self.client.post("/api/v2/pathways", headers=headers, json=request)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["plan"]["selected_software_skill_ids"], [selected])
        self.assertEqual(response.json()["plan"]["fit_score"], role["fit_score"])
        invalid = self.client.post("/api/v2/pathways", headers=headers,
            json={**request, "selected_software_skill_ids": ["skill:invented"]})
        self.assertEqual(invalid.status_code, 422)


class TestDocumentFormats(unittest.TestCase):
    def test_docx_reads_tables_and_rejects_malformed_files(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Fictional graduate")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Skills: Python"
        table.cell(0, 1).text = "SQL, Tableau"
        content = BytesIO()
        document.save(content)
        text = "\n".join(parse_cv(content.getvalue(), "cv.docx").pages)
        self.assertIn("Python", text)
        self.assertIn("SQL", text)
        with self.assertRaises(CVParseError):
            parse_cv(b"not a word document", "cv.docx")

    def test_image_and_scanned_pdf_run_real_local_ocr(self):
        from PIL import Image, ImageDraw, ImageFont

        image = Image.new("RGB", (1200, 400), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.load_default(size=38)
        draw.text((45, 55), "FICTIONAL CV", fill="black", font=font)
        draw.text((45, 140), "Skills: Python, SQL, Tableau", fill="black", font=font)
        for kind, suffix in [("PNG", "png"), ("JPEG", "jpg"), ("PDF", "pdf")]:
            with self.subTest(format=kind):
                content = BytesIO()
                image.save(content, format=kind)
                parsed = parse_cv(content.getvalue(), "cv." + suffix)
                text = "\n".join(parsed.pages)
                self.assertIn("Python", text)
                self.assertIn("SQL", text)


if __name__ == "__main__":
    unittest.main()
