"""Create fictional cross-domain CVs for manual upload and parser checks."""

import json
import textwrap
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from src.profiles.samples import SAMPLE_PATH, sample_cv_pdf

DESTINATION = SAMPLE_PATH.parent / "domain_examples"
EXAMPLES = [
    {
        "id": "data_analytics", "domain": "Data analytics", "name": "Nadia Lim",
        "title": "Data Analyst", "primary_format": "pdf",
        "education": "Bachelor of Computing | Fictional Harbour Institute | 2023",
        "experience": "Data Analyst | Fictional Northstar Retail | 2023 - 2026",
        "bullets": [
            "Used Python and SQL to clean synthetic sales records and validate monthly totals.",
            "Built Tableau and Microsoft Power BI dashboards to compare regional performance.",
            "Used Microsoft Excel to reconcile report totals and document data-quality checks.",
            "Applied Critical Thinking and Mathematics to investigate inconsistent measurements.",
            "Demonstrated Writing and Reading Comprehension through clear reporting notes.",
        ],
        "project": "Retail Reporting Pipeline",
        "project_text": "Delivered a documented dashboard with repeatable SQL queries and Python checks.",
        "skills": "Python, SQL, Tableau, Microsoft Power BI, Microsoft Excel, Critical Thinking, Mathematics, Writing, Reading Comprehension",
    },
    {
        "id": "web_development", "domain": "Web development", "name": "Daniel Arif",
        "title": "Frontend Web Developer", "primary_format": "docx",
        "education": "Diploma in Software Development | Fictional Bay Academy | 2023",
        "experience": "Web Developer | Fictional Lantern Studio | 2023 - 2026",
        "bullets": [
            "Built accessible interfaces with React, JavaScript, HTML and CSS for a fictional portal.",
            "Used Node.js to implement API endpoints and Git to review versioned changes.",
            "Applied Programming and Complex Problem Solving to fix form validation defects.",
            "Used Critical Thinking to investigate errors and Writing to explain release changes.",
        ],
        "project": "Community Booking Portal",
        "project_text": "Implemented responsive forms, keyboard navigation and automated interface checks.",
        "skills": "React, JavaScript, HTML, CSS, Node.js, Git, Programming, Complex Problem Solving, Critical Thinking, Writing",
    },
    {
        "id": "healthcare_nursing", "domain": "Healthcare / nursing", "name": "Sara Menon",
        "title": "Registered Nurse - Fictional Training Profile", "primary_format": "txt",
        "education": "Bachelor of Nursing | Fictional Meadow College | 2022",
        "experience": "Nurse | Fictional Meadow Training Hospital | 2022 - 2026",
        "bullets": [
            "Used Epic Systems and MEDITECH software to document simulated patient observations.",
            "Used PointClickCare healthcare software to record fictional handover notes.",
            "Practised Active Listening, Social Perceptiveness and Service Orientation in simulation.",
            "Applied Monitoring and Coordination during supervised clinical training exercises.",
            "Demonstrated Speaking, Critical Thinking and Judgment and Decision Making in handovers.",
        ],
        "project": "Simulated Ward Handover Improvement",
        "project_text": "Created a checklist to communicate observations and escalation responsibilities.",
        "skills": "Epic Systems, MEDITECH software, PointClickCare healthcare software, Active Listening, Social Perceptiveness, Service Orientation, Monitoring, Coordination, Speaking, Critical Thinking, Judgment and Decision Making",
    },
    {
        "id": "finance_accounting", "domain": "Finance / accounting", "name": "Marcus Lee",
        "title": "Accounts Executive", "primary_format": "png",
        "education": "Bachelor of Accounting | Fictional Meridian College | 2022",
        "experience": "Accounts Executive | Fictional Cedar Services | 2022 - 2026",
        "bullets": [
            "Used Intuit QuickBooks and Sage 50 Accounting to prepare fictional monthly ledgers.",
            "Used SAP ERP Financials and Microsoft Excel to reconcile sample invoices and balances.",
            "Applied Mathematics and Critical Thinking to investigate duplicate payment entries.",
            "Used Monitoring and Writing to maintain audit checklists and explain adjustments.",
            "Applied Reading Comprehension to follow documented reconciliation procedures.",
        ],
        "project": "Month-End Reconciliation Pack",
        "project_text": "Prepared a traceable spreadsheet and supporting notes for a simulated review.",
        "skills": "Intuit QuickBooks, Sage 50 Accounting, SAP ERP Financials, Microsoft Excel, Mathematics, Critical Thinking, Monitoring, Writing, Reading Comprehension",
    },
    {
        "id": "mechanical_engineering", "domain": "Mechanical engineering", "name": "Amir Roslan",
        "title": "Mechanical Design Engineer", "primary_format": "jpg",
        "education": "Bachelor of Mechanical Engineering | Fictional Crest University | 2021",
        "experience": "Design Engineer | Fictional Crest Manufacturing | 2021 - 2026",
        "bullets": [
            "Used Dassault Systemes SolidWorks and Autodesk AutoCAD Mechanical to model a fixture.",
            "Used PTC Creo Parametric to revise prototype drawings and document dimensional checks.",
            "Used The MathWorks MATLAB to compare simulated loads and review calculation outputs.",
            "Applied Mathematics, Science and Complex Problem Solving during prototype evaluation.",
            "Demonstrated Operations Analysis, Critical Thinking and Writing in design reviews.",
        ],
        "project": "Lightweight Test Fixture",
        "project_text": "Produced drawings, a calculation sheet and a supervised prototype test report.",
        "skills": "Dassault Systemes SolidWorks, Autodesk AutoCAD Mechanical, PTC Creo Parametric, The MathWorks MATLAB, Mathematics, Science, Complex Problem Solving, Operations Analysis, Critical Thinking, Writing",
    },
]


def cv_text(example):
    return "\n".join([
        example["name"].upper(), example["title"],
        example["id"] + "@example.com | Kuala Lumpur, Malaysia",
        "FICTIONAL TEST CV - all people, employers and qualifications are invented.",
        "", "PROFILE", "Practical experience in " + example["domain"].lower() + ".",
        "", "EDUCATION", example["education"], "", "EXPERIENCE",
        example["experience"], *["- " + line for line in example["bullets"]],
        "", "PROJECT", example["project"], example["project_text"],
        "", "SKILLS", example["skills"], "",
    ])


def pdf_bytes(text):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    fonts = DictionaryObject()
    for key, font in (("/F1", "/Helvetica"), ("/F2", "/Helvetica-Bold")):
        fonts[NameObject(key)] = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject(font)})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): fonts})
    commands, y = [], 793
    for index, line in enumerate(text.splitlines()):
        if not line:
            y -= 9
            continue
        bold = index < 2 or (line.isupper() and len(line) < 24)
        size = 21 if index == 0 else 11 if bold else 9.5
        for wrapped in textwrap.wrap(line, width=87 if index else 40):
            escaped = wrapped.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            color = "0.08 0.23 0.31" if bold else "0.16 0.18 0.2"
            commands.append(f"{color} rg BT /{'F2' if bold else 'F1'} {size} Tf 1 0 0 1 43 {y} Tm ({escaped}) Tj ET")
            y -= 15 if index else 29
    if y < 40:
        raise ValueError("CV exceeds one page")
    stream = DecodedStreamObject()
    stream.set_data("\n".join(commands).encode("ascii"))
    page.replace_contents(stream)
    writer.add_metadata({"/Title": text.splitlines()[0] + " - Fictional CV"})
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def write_docx(text, path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.65)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10)
    style.paragraph_format.space_after = Pt(4)
    for index, line in enumerate(text.splitlines()):
        if not line:
            continue
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(line)
        if index < 2 or (line.isupper() and len(line) < 24):
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("143B4F")
            run.font.size = Pt(21 if index == 0 else 11)
    doc.save(path)


def write_image(text, path):
    image = Image.new("RGB", (1400, 1900), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    bold_path = Path("C:/Windows/Fonts/arialbd.ttf")
    def font(size, bold=False):
        chosen = bold_path if bold else font_path
        return ImageFont.truetype(str(chosen), size) if chosen.is_file() else ImageFont.load_default(size=size)
    y = 65
    for index, line in enumerate(text.splitlines()):
        if not line:
            y += 17
            continue
        bold = index < 2 or (line.isupper() and len(line) < 24)
        size = 48 if index == 0 else 30 if bold else 25
        for wrapped in textwrap.wrap(line, width=88 if index else 36):
            draw.text((65, y), wrapped, font=font(size, bold), fill="#143b4f" if bold else "#232b31")
            y += 35 if index else 65
    if y > image.height - 50:
        raise ValueError("CV image exceeds page")
    image.save(path, quality=95, dpi=(170, 170))


def main():
    DESTINATION.mkdir(parents=True, exist_ok=True)
    manifest = []
    for example in EXAMPLES:
        text = cv_text(example)
        stem = DESTINATION / example["id"]
        stem.with_suffix(".txt").write_text(text, encoding="utf-8")
        stem.with_suffix(".pdf").write_bytes(pdf_bytes(text))
        write_docx(text, stem.with_suffix(".docx"))
        formats = ["pdf", "docx", "txt"]
        if example["primary_format"] in {"png", "jpg"}:
            write_image(text, stem.with_suffix("." + example["primary_format"]))
            formats.append(example["primary_format"])
        manifest.append({"id": example["id"], "domain": example["domain"],
            "primary_file": example["id"] + "." + example["primary_format"],
            "files": [example["id"] + "." + extension for extension in formats]})
    SAMPLE_PATH.with_suffix(".pdf").write_bytes(sample_cv_pdf())
    (DESTINATION / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(f"Created {sum(len(row['files']) for row in manifest)} CV examples in {DESTINATION}")
    print(f"Existing sample PDF: {SAMPLE_PATH.with_suffix('.pdf')}")


if __name__ == "__main__":
    main()
