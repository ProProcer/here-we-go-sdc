"""Create different visual compositions from the same fictional CV content."""

import json
import textwrap
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from scripts.create_cv_examples import DESTINATION, EXAMPLES, cv_text, pdf_bytes, write_docx

OUTPUT = DESTINATION.parent / "layout_examples"


def sidebar_pdf(example):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    fonts = DictionaryObject()
    for key, face in (("/F1", "/Helvetica"), ("/F2", "/Helvetica-Bold")):
        fonts[NameObject(key)] = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject(face)})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): fonts})
    commands = ["0.08 0.23 0.31 rg 0 0 205 842 re f"]
    def draw(text, x, y, *, width, size=10, bold=False, white=False):
        for line in textwrap.wrap(text, width=width):
            value = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            color = "1 1 1" if white else "0.08 0.23 0.31"
            commands.append(f"{color} rg BT /{'F2' if bold else 'F1'} {size} Tf 1 0 0 1 {x} {y} Tm ({value}) Tj ET")
            y -= size + 6
        return y
    y = draw(example["name"].upper(), 22, 790, width=17, size=20, bold=True, white=True) - 12
    y = draw(example["title"], 22, y, width=25, size=11, bold=True, white=True) - 20
    y = draw("SKILLS", 22, y, width=24, bold=True, white=True) - 5
    for skill in example["skills"].split(", "):
        y = draw(skill, 22, y, width=30, size=9, white=True) - 3
    y = draw("EDUCATION", 22, y - 15, width=24, bold=True, white=True)
    draw(example["education"], 22, y - 7, width=29, size=9, white=True)
    y = 790
    for title, content in [
        ("FICTIONAL TEST CV", "All names, employers and qualifications are invented."),
        ("PROFILE", "Practical experience in " + example["domain"].lower() + "."),
        ("EXPERIENCE", example["experience"]),
    ]:
        y = draw(title, 230, y, width=47, size=12, bold=True) - 6
        y = draw(content, 230, y, width=59, size=9) - 22
    for bullet in example["bullets"]:
        y = draw("- " + bullet, 230, y, width=59, size=9) - 12
    y = draw("PROJECT", 230, y - 10, width=47, size=12, bold=True) - 7
    y = draw(example["project"], 230, y, width=59, size=10, bold=True) - 7
    draw(example["project_text"], 230, y, width=59, size=9)
    stream = DecodedStreamObject()
    stream.set_data("\n".join(commands).encode("ascii"))
    page.replace_contents(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def table_docx(example, path):
    document = Document()
    document.sections[0].top_margin = Inches(0.6)
    document.sections[0].bottom_margin = Inches(0.6)
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading(example["name"], 0)
    document.add_paragraph(example["title"])
    document.add_paragraph("FICTIONAL TEST CV - all qualifications and experience are invented.")
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    sidebar, main = table.rows[0].cells
    sidebar.text = "SKILLS\n" + example["skills"].replace(", ", "\n") + "\n\nEDUCATION\n" + example["education"]
    main.text = "PROFILE\nPractical experience in " + example["domain"].lower() + ".\n\nEXPERIENCE\n" + example["experience"]
    for line in example["bullets"]:
        main.add_paragraph(line)
    main.add_paragraph("PROJECT\n" + example["project"] + "\n" + example["project_text"])
    document.save(path)


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    examples = {example["id"]: example for example in EXAMPLES}
    manifest = []
    def record(filename, source, composition):
        manifest.append({"file": filename, "source": source + ".txt", "composition": composition})
    data = examples["data_analytics"]
    (OUTPUT / "analytics_ats_single_column.pdf").write_bytes(pdf_bytes(cv_text(data)))
    record("analytics_ats_single_column.pdf", "data_analytics", "ATS-friendly single column with standard headings; text-based PDF")
    (OUTPUT / "analytics_sidebar.pdf").write_bytes(sidebar_pdf(data))
    record("analytics_sidebar.pdf", "data_analytics", "Two columns, dark skills sidebar and separate experience column")
    text = cv_text(data)
    header, rest = text.split("\nPROFILE\n", 1)
    profile, rest = rest.split("\nEDUCATION\n", 1)
    education, rest = rest.split("\nEXPERIENCE\n", 1)
    experience, rest = rest.split("\nPROJECT\n", 1)
    project, skills = rest.split("\nSKILLS\n", 1)
    functional = "\n".join([header, "SKILLS", skills, "PROJECT", project, "PROFILE", profile,
        "EXPERIENCE", experience, "EDUCATION", education])
    (OUTPUT / "analytics_functional.pdf").write_bytes(pdf_bytes(functional))
    record("analytics_functional.pdf", "data_analytics", "Functional composition: competencies and project precede work history")
    write_docx(text, OUTPUT / "analytics_chronological.docx")
    record("analytics_chronological.docx", "data_analytics", "Chronological, single-column Word document")
    table_docx(examples["healthcare_nursing"], OUTPUT / "nursing_table_layout.docx")
    record("nursing_table_layout.docx", "healthcare_nursing", "Two-column Word table with skills and education in a sidebar")
    web = cv_text(examples["web_development"])
    first, second = web.split("EXPERIENCE\n", 1)
    writer = PdfWriter()
    for part in (first, "DANIEL ARIF\nFICTIONAL CV - CONTINUED\n\nEXPERIENCE\n" + second):
        writer.append(PdfReader(BytesIO(pdf_bytes(part))))
    with (OUTPUT / "web_two_page.pdf").open("wb") as output:
        writer.write(output)
    record("web_two_page.pdf", "web_development", "Two-page CV with experience and skills continuing on page two")
    with Image.open(DESTINATION / "finance_accounting.png") as image:
        image.convert("RGB").save(OUTPUT / "accounting_scanned.pdf", "PDF", resolution=150)
    record("accounting_scanned.pdf", "finance_accounting", "Scanned image-only PDF; no embedded text, requires OCR")
    table_docx(examples["mechanical_engineering"], OUTPUT / "engineering_styled.docx")
    document = Document(OUTPUT / "engineering_styled.docx")
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string("173D57")
    document.save(OUTPUT / "engineering_styled.docx")
    record("engineering_styled.docx", "mechanical_engineering", "Styled Word CV with colored header and shaded two-column table")
    (OUTPUT / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(f"Created {len(manifest)} distinct CV layouts in {OUTPUT}")


if __name__ == "__main__":
    main()
